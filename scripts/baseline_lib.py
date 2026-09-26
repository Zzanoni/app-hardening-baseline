"""Shared helpers for the baseline workbook scripts (migrate, export, check, upgrade).

Everything that more than one script needs lives here: contract loading, the condition language
(translation to Excel formulas, human-readable labels and a Python evaluator), deterministic
saving, LibreOffice recalculation and data-validation lookup.
"""

import datetime
import io
import os
import re
import shutil
import subprocess
import tempfile
import zipfile

import yaml
from openpyxl import load_workbook
from openpyxl.utils import range_boundaries

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
TEMPLATE = os.path.join(ROOT, "templates", "ASVS_Master_Hardening_Template.xlsx")
CONTRACT_DIR = os.path.join(ROOT, "contract")
CONTROLS_JSON = os.path.join(CONTRACT_DIR, "controls.json")

SHEET_CF = "Characterization Form"
SHEET_EXT = "Extended Characterization"
SHEET_MT = "Master Template"
SHEET_CC = "Custom Controls"
SHEET_ARCH = "Archetypes"
SHEET_LISTS = "_lists"
SHEET_META = "_meta"

# Fixed package timestamps so that saving the same content always yields the same bytes.
FIXED_ZIP_DATE = (1980, 1, 1, 0, 0, 0)
FIXED_MODIFIED = "2026-09-26T00:00:00Z"

ERROR_VALUES = ("#NAME?", "#VALUE!", "#REF!", "#DIV/0!", "#N/A", "#NUM!", "#NULL!")


# --------------------------------------------------------------------------- contract


def load_yaml(name):
    with open(os.path.join(CONTRACT_DIR, name), encoding="utf-8") as fh:
        return yaml.safe_load(fh)


def load_contract():
    """Return (characterization, archetypes, workbook_contract) as plain dicts."""
    return (
        load_yaml("characterization.yml"),
        load_yaml("archetypes.yml"),
        load_yaml("workbook-contract.yml"),
    )


def keys_for_sheet(characterization, sheet):
    return [k for k in characterization["keys"] if k["sheet"] == sheet]


def options_label(key):
    if key.get("options_label"):
        return key["options_label"]
    kind = key["type"]
    if kind == "text":
        return "Free text"
    if kind == "long_text":
        return "Free text (2-3 lines)"
    if kind == "date":
        return "Date (YYYY-MM-DD)"
    options = key["options"]
    sep = "; " if any("," in o for o in options) else ", "
    return sep.join(options)


# --------------------------------------------------------------------- condition language


def _excel_str(value):
    return '"' + str(value).replace('"', '""') + '"'


def condition_to_excel(cond, ref, special=None):
    """Translate a contract condition into an Excel-2007 boolean expression.

    ref(key) -> (value_expr, guard_expr_or_None). A guard is the key's own applicability; when
    present the comparison only holds if the key is applicable (non-applicable answers count as
    blank). special maps pseudo-keys (e.g. status_any) to callables value -> expression.
    """
    special = special or {}
    parts = []
    for op, arg in cond.items():
        if op == "not":
            parts.append("NOT(%s)" % condition_to_excel(arg, ref, special))
        elif op == "all_of":
            parts.append("AND(%s)" % ",".join(condition_to_excel(c, ref, special) for c in arg))
        elif op == "any_of":
            for key, values in arg.items():
                value_expr, guard = ref(key)
                expr = "OR(%s)" % ",".join("%s=%s" % (value_expr, _excel_str(v)) for v in values)
                parts.append("AND(%s,%s)" % (guard, expr) if guard else expr)
        elif op in special:
            parts.append(special[op](arg))
        else:
            value_expr, guard = ref(op)
            expr = "%s=%s" % (value_expr, _excel_str(arg))
            parts.append("AND(%s,%s)" % (guard, expr) if guard else expr)
    return parts[0] if len(parts) == 1 else "AND(%s)" % ",".join(parts)


def condition_label(cond):
    """Human-readable rendering of a condition, for the 'Shown when' column."""
    parts = []
    for op, arg in cond.items():
        if op == "not":
            if len(arg) == 1 and next(iter(arg)) not in ("not", "all_of", "any_of"):
                key, value = next(iter(arg.items()))
                parts.append("%s ≠ %s" % (key, value))
            else:
                parts.append("not (%s)" % condition_label(arg))
        elif op == "all_of":
            parts.append(" and ".join(condition_label(c) for c in arg))
        elif op == "any_of":
            for key, values in arg.items():
                parts.append("%s is one of: %s" % (key, "; ".join(values)))
        else:
            parts.append("%s = %s" % (op, arg))
    return " and ".join(parts)


def evaluate_condition(cond, value_of, special=None):
    """Python evaluator with the same semantics as condition_to_excel.

    value_of(key) must already return "" for keys whose row is not applicable.
    """
    special = special or {}
    results = []
    for op, arg in cond.items():
        if op == "not":
            results.append(not evaluate_condition(arg, value_of, special))
        elif op == "all_of":
            results.append(all(evaluate_condition(c, value_of, special) for c in arg))
        elif op == "any_of":
            results.append(all(value_of(k) in values for k, values in arg.items()))
        elif op in special:
            results.append(special[op](arg))
        else:
            results.append(value_of(op) == arg)
    return all(results)


def referenced_keys(cond):
    found = []
    for op, arg in cond.items():
        if op == "not":
            found += referenced_keys(arg)
        elif op == "all_of":
            for c in arg:
                found += referenced_keys(c)
        elif op == "any_of":
            found += list(arg)
        elif op != "status_any":
            found.append(op)
    return found


# ------------------------------------------------------------------------- workbook I/O


def save_normalized(wb, path):
    """Save with fixed package timestamps, so identical content gives identical bytes."""
    buf = io.BytesIO()
    wb.save(buf)
    normalize_package(buf.getvalue(), path)


def normalize_package(data, path):
    src = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            content = src.read(info.filename)
            if info.filename == "docProps/core.xml":
                content = re.sub(
                    rb"(<dcterms:modified[^>]*>)[^<]*(</dcterms:modified>)",
                    rb"\g<1>" + FIXED_MODIFIED.encode() + rb"\g<2>",
                    content,
                )
            zi = zipfile.ZipInfo(info.filename, date_time=FIXED_ZIP_DATE)
            zi.compress_type = zipfile.ZIP_DEFLATED
            zi.external_attr = 0o600 << 16
            dst.writestr(zi, content)
    with open(path, "wb") as fh:
        fh.write(out.getvalue())


def defined_name_ref(wb, name):
    dn = wb.defined_names.get(name)
    return dn.attr_text if dn is not None else None


def parse_ref(ref):
    """"'Sheet'!$D$5" -> ("Sheet", "D5"); ranges keep their colon ("B26:E28")."""
    sheet, cells = ref.rsplit("!", 1)
    return sheet.strip("'").replace("''", "'"), cells.replace("$", "")


def named_values(wb, name):
    """Values of a named range, as a flat list (column-major for single-column ranges)."""
    sheet, cells = parse_ref(defined_name_ref(wb, name))
    ws = wb[sheet]
    min_col, min_row, max_col, max_row = range_boundaries(cells)
    values = []
    for row in ws.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col):
        values.extend(c.value for c in row)
    return values


def validation_for(ws, coordinate):
    for dv in ws.data_validations.dataValidation:
        if coordinate in dv.sqref:
            return dv
    return None


def validation_options(wb, dv):
    """Resolve a list data validation to its options (inline list or named range)."""
    formula = (dv.formula1 or "").strip()
    if formula.startswith('"'):
        return formula.strip('"').split(",")
    name = formula.lstrip("=")
    return [v for v in named_values(wb, name) if v is not None]


# ------------------------------------------------------------------------ recalculation

_LO_CANDIDATES = (
    "soffice",
    "libreoffice",
    "/Applications/LibreOffice.app/Contents/MacOS/soffice",
)

_LO_PROFILE = """<?xml version="1.0" encoding="UTF-8"?>
<oor:items xmlns:oor="http://openoffice.org/2001/registry"
           xmlns:xs="http://www.w3.org/2001/XMLSchema"
           xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">
  <item oor:path="/org.openoffice.Office.Calc/Formula/Load">
    <prop oor:name="OOXMLRecalcMode" oor:op="fuse"><value>0</value></prop>
  </item>
</oor:items>
"""


def find_soffice():
    for candidate in _LO_CANDIDATES:
        found = shutil.which(candidate) or (candidate if os.path.isfile(candidate) else None)
        if found:
            return found
    return None


def recalculate(path, soffice=None):
    """Recalculate a workbook with LibreOffice headless; return the path of the result.

    Uses a throwaway LibreOffice profile configured to always recalculate OOXML files on load.
    The caller owns the returned file's directory (a fresh temp dir).
    """
    soffice = soffice or find_soffice()
    if not soffice:
        raise RuntimeError("LibreOffice (soffice) not found")
    work = tempfile.mkdtemp(prefix="baseline-recalc-")
    profile = os.path.join(work, "profile")
    os.makedirs(os.path.join(profile, "user"))
    with open(os.path.join(profile, "user", "registrymodifications.xcu"), "w") as fh:
        fh.write(_LO_PROFILE)
    src = os.path.join(work, "in", os.path.basename(path))
    os.makedirs(os.path.dirname(src))
    shutil.copy(path, src)
    outdir = os.path.join(work, "out")
    subprocess.run(
        [
            soffice,
            "-env:UserInstallation=file://" + profile,
            "--headless",
            "--norestore",
            "--convert-to",
            "xlsx:Calc MS Excel 2007 XML",
            "--outdir",
            outdir,
            src,
        ],
        check=True,
        stdout=subprocess.DEVNULL,
        stderr=subprocess.DEVNULL,
        timeout=300,
    )
    result = os.path.join(outdir, os.path.basename(path))
    if not os.path.exists(result):
        raise RuntimeError("LibreOffice produced no output for " + path)
    return result


def formula_errors(path):
    """List (sheet, cell, value) for every cell whose computed value is an error."""
    wb = load_workbook(path, data_only=True)
    errors = []
    for ws in wb:
        for row in ws.iter_rows():
            for cell in row:
                v = cell.value
                if isinstance(v, str) and (v in ERROR_VALUES or v.startswith("Err:")):
                    errors.append((ws.title, cell.coordinate, v))
    return errors


