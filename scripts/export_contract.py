#!/usr/bin/env python3
"""Export the control catalog of the master workbook to contract/controls.json.

One entry per ASVS requirement (Master Template) and per non-empty Custom Controls row, with
the fields the threat engine needs to link threats to controls. Output is deterministic (no
timestamps, stable order) so check_contract.py can detect drift by regenerating and comparing.

Usage: python scripts/export_contract.py [--workbook PATH] [--out PATH]
"""

import argparse
import json

from openpyxl import load_workbook

import baseline_lib as lib


def _text(value):
    if value is None:
        return None
    value = str(value).strip()
    return value or None


def _formula(value):
    value = _text(value)
    return value[1:] if value and value.startswith("=") else value


def build_catalog(workbook_path):
    wb = load_workbook(workbook_path)
    _, _, contract = lib.load_contract()
    sheets = {s["name"]: s for s in contract["sheets"]}
    examples = set(contract["example_control_ids"])
    controls = []
    for sheet_name, source in ((lib.SHEET_MT, "asvs"), (lib.SHEET_CC, "custom")):
        spec = sheets[sheet_name]
        ws = wb[sheet_name]
        for r in range(spec["first_data_row"], spec["last_data_row"] + 1):
            control_id = _text(ws["A%d" % r].value)
            if not control_id:
                continue
            level = ws["D%d" % r].value
            controls.append({
                "id": control_id,
                "source": source,
                "source_sheet": sheet_name,
                "chapter" if source == "asvs" else "category": _text(ws["B%d" % r].value),
                "section" if source == "asvs" else "sub_category": _text(ws["C%d" % r].value),
                "level": int(level) if level not in (None, "") else None,
                "requirement": _text(ws["E%d" % r].value),
                "applicability_condition": _text(ws["F%d" % r].value),
                "applicability_formula": _formula(ws["G%d" % r].value),
                "control_type": _text(ws["I%d" % r].value),
                "nature": _text(ws["J%d" % r].value),
                "verification_stage": _text(ws["K%d" % r].value),
                "requirement_impact": _text(ws["O%d" % r].value),
                "is_example": control_id in examples,
            })
    meta = {name: lib.named_values(wb, name)[0] for name in contract["meta_named_ranges"]}
    return {
        "contract_version": meta["CONTRACT_VERSION"],
        "template_version": meta["TEMPLATE_VERSION"],
        "asvs_version": meta["ASVS_VERSION"],
        "generated_by": "scripts/export_contract.py — do not edit by hand",
        "counts": {
            "asvs": sum(1 for c in controls if c["source"] == "asvs"),
            "custom": sum(1 for c in controls if c["source"] == "custom"),
        },
        "controls": controls,
    }


def render(catalog):
    return json.dumps(catalog, indent=2, ensure_ascii=False) + "\n"


def main():
    parser = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    parser.add_argument("--workbook", default=lib.TEMPLATE)
    parser.add_argument("--out", default=lib.CONTROLS_JSON)
    args = parser.parse_args()
    catalog = build_catalog(args.workbook)
    with open(args.out, "w", encoding="utf-8") as fh:
        fh.write(render(catalog))
    print("Wrote %s (%d ASVS + %d custom controls)" % (args.out, catalog["counts"]["asvs"], catalog["counts"]["custom"]))


if __name__ == "__main__":
    main()
