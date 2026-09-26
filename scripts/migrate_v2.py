#!/usr/bin/env python3
"""Migrate the master workbook from template v1 to v2.0.0 (threat modeling integration).

Builds everything v2 adds from the contract files in contract/:
  - `Extended Characterization` sheet (from characterization.yml)
  - `Archetypes` reference sheet (from archetypes.yml)
  - hidden `_lists` (dropdown sources) and `_meta` (versions) sheets
  - `Remediation Owner` column (Q) on `Master Template` and `Custom Controls`
  - `Custom Controls` extended to rows 5-204, with a TMX-EXAMPLE-001 worked example
  - updated `Instructions`
  - APP_NAME named range (was missing in v1)

Idempotent: every v2 part is rebuilt from scratch on each run and the file is saved with fixed
package timestamps, so running it on its own output produces a byte-identical file. v1 content
(named ranges, answer cells, computed cells, matrix, Master Template A-P, Custom Controls
formulas) is snapshotted before and after; the script refuses to write if any of it changed.

Usage: python scripts/migrate_v2.py [--in PATH] [--out PATH]   (default: the template, in place)
"""

import argparse
import copy
import sys

from openpyxl import load_workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.formatting.formatting import ConditionalFormattingList
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.cell_range import MultiCellRange
from openpyxl.worksheet.datavalidation import DataValidation

import baseline_lib as lib

MT_LAST_ROW = 349
CC_FIRST_ROW = 5
CC_V1_LAST_ROW = 24
CC_LAST_ROW = 204
TMX_EXAMPLE_ROW = 6

V1_KEYS = [
    "APP_NAME", "TARGET_LEVEL", "API_EXPOSED", "SPA_FRONTEND", "SESSION_STATEFUL", "FILE_HANDLING",
    "OAUTH_OIDC_USED", "DATA_SENSITIVITY", "WEBRTC_USED", "EXPOSURE", "BUSINESS_IMPACT",
]
V1_NAMES = [
    "API_EXPOSED", "FILE_HANDLING", "OAUTH_OIDC_USED", "SESSION_STATEFUL", "SPA_FRONTEND",
    "TARGET_LEVEL", "WEBRTC_USED", "DATA_SENSITIVITY", "EXPOSURE", "BUSINESS_IMPACT",
    "TARGET_LEVEL_NUM", "DATA_SENSITIVITY_SCORE", "EXPOSURE_SCORE", "BUSINESS_IMPACT_SCORE",
    "APP_RISK_SCORE", "APP_RISK_TIER", "CRIT_MATRIX", "CRIT_MATRIX_ROWS", "CRIT_MATRIX_COLS",
]

SECTION_FILL = PatternFill("solid", fgColor="FFD9D9D9")
SECTION_FONT = Font(name="Calibri", size=10, bold=True, color="FF262626")
RED_FILL = PatternFill("solid", fgColor="FFFFC7CE", bgColor="FFFFC7CE")
AMBER_FILL = PatternFill("solid", fgColor="FFFFEB9C", bgColor="FFFFEB9C")
GREY_FILL = PatternFill("solid", fgColor="FFF2F2F2", bgColor="FFF2F2F2")

REMEDIATION_HELP = (
    "Remediation Owner — who can close this gap. Fill only for rows with Coverage Status "
    "'Partial coverage' or 'Gap'.\n"
    "Internal dev: code change by our developers.\n"
    "Configuration: a setting in the product, platform or infrastructure.\n"
    "Vendor: only the vendor can fix it (product change or patch).\n"
    "Compensating: fixed outside the product (WAF, segmentation, PAM, monitoring).\n"
    "Risk acceptance: formally accepted instead of fixed.\n"
    "For COTS, most application-level gaps are Vendor or Configuration."
)


# ------------------------------------------------------------------------------ helpers


def copy_style(src, dst):
    dst._style = copy.copy(src._style)


def set_name(wb, name, ref):
    if name in wb.defined_names:
        del wb.defined_names[name]
    wb.defined_names[name] = DefinedName(name, attr_text=ref)


def abs_ref(sheet, cells):
    quoted = "'%s'" % sheet.replace("'", "''")
    if ":" in cells:
        a, b = cells.split(":")
        return "%s!%s:%s" % (quoted, _abs(a), _abs(b))
    return "%s!%s" % (quoted, _abs(cells))


def _abs(cell):
    col = "".join(ch for ch in cell if ch.isalpha())
    row = "".join(ch for ch in cell if ch.isdigit())
    return "$%s$%s" % (col, row)


def recreate_sheet(wb, title, index):
    if title in wb.sheetnames:
        wb.remove(wb[title])
    ws = wb.create_sheet(title, index)
    ws.sheet_view.showGridLines = False
    return ws


def remove_validations_on(ws, coordinate):
    for dv in list(ws.data_validations.dataValidation):
        if coordinate in dv.sqref:
            ws.data_validations.dataValidation.remove(dv)


def list_validation(formula, sqref):
    dv = DataValidation(type="list", formula1=formula, allow_blank=True)
    dv.prompt = "Choose an option."
    dv.showInputMessage = True
    dv.error = "Choose one of the listed options."
    dv.showErrorMessage = True
    dv.add(sqref)
    return dv


# ----------------------------------------------------------------------- v1 invariants


def snapshot_v1(wb):
    """Everything v1 that v2 must leave untouched."""
    snap = {"names": {n: lib.defined_name_ref(wb, n) for n in V1_NAMES}}
    cf = wb[lib.SHEET_CF]
    snap["cf_cells"] = {c.coordinate: c.value for row in cf.iter_rows() for c in row}
    snap["cf_dv"] = sorted((str(dv.sqref), dv.formula1) for dv in cf.data_validations.dataValidation)
    mt = wb[lib.SHEET_MT]
    snap["mt_cells"] = {
        c.coordinate: c.value
        for row in mt.iter_rows(min_row=1, max_row=MT_LAST_ROW, max_col=16)
        for c in row
    }
    snap["mt_hidden_g"] = mt.column_dimensions["G"].hidden
    snap["mt_freeze"] = mt.freeze_panes
    snap["mt_dv_l"] = [dv.formula1 for dv in mt.data_validations.dataValidation if "L5" in dv.sqref]
    snap["mt_cf"] = sorted(
        (str(cf_range.sqref), r.operator, tuple(r.formula))
        for cf_range in mt.conditional_formatting
        for r in cf_range.rules
    )
    cc = wb[lib.SHEET_CC]
    cc_formulas = {}
    for r in range(CC_FIRST_ROW, CC_V1_LAST_ROW + 1):
        for col in ("H", "P") + (("G",) if r != TMX_EXAMPLE_ROW else ()):
            cc_formulas[col + str(r)] = cc[col + str(r)].value
    snap["cc_formulas"] = cc_formulas
    snap["cc_example"] = [cc.cell(CC_FIRST_ROW, c).value for c in range(1, 17)]
    snap["cc_headers"] = [cc.cell(4, c).value for c in range(1, 17)]
    return snap


def check_v1_layout(wb, characterization):
    missing = [s for s in (lib.SHEET_CF, lib.SHEET_MT, lib.SHEET_CC, "Instructions") if s not in wb.sheetnames]
    if missing:
        sys.exit("Not a v1/v2 workbook, missing sheets: %s" % missing)
    cf = wb[lib.SHEET_CF]
    found = [cf.cell(r, 1).value for r in range(5, 16)]
    if found != V1_KEYS:
        sys.exit("Characterization Form rows 5-15 do not hold the v1 keys: %s" % found)
    contract_v1 = [k["key"] for k in lib.keys_for_sheet(characterization, lib.SHEET_CF)]
    if contract_v1 != V1_KEYS:
        sys.exit("characterization.yml v1 keys differ from the workbook: %s" % contract_v1)


# --------------------------------------------------------------------------- builders


def build_lists(wb, characterization, archetypes):
    lists = {}
    for key in characterization["keys"]:
        name = key.get("list")
        if not name:
            continue
        if name in lists and lists[name] != key["options"]:
            sys.exit("List %s is used with different options (%s)" % (name, key["key"]))
        lists[name] = key["options"]
    lists["LIST_ARCHETYPES"] = [a["name"] for a in archetypes["archetypes"]]

    ws = recreate_sheet(wb, lib.SHEET_LISTS, len(wb.sheetnames))
    ws.sheet_state = "hidden"
    for col, name in enumerate(sorted(lists), start=1):
        letter = get_column_letter(col)
        ws.cell(1, col, name).font = Font(bold=True)
        for i, value in enumerate(lists[name], start=2):
            ws.cell(i, col, value)
        ws.column_dimensions[letter].width = 40
        set_name(wb, name, abs_ref(lib.SHEET_LISTS, "%s2:%s%d" % (letter, letter, len(lists[name]) + 1)))
    return lists


def build_meta(wb, contract):
    ws = recreate_sheet(wb, lib.SHEET_META, len(wb.sheetnames))
    ws.sheet_state = "hidden"
    ws["A1"], ws["B1"] = "Name", "Value"
    ws["A1"].font = ws["B1"].font = Font(bold=True)
    for row, (name, field) in enumerate(sorted(contract["meta_named_ranges"].items()), start=2):
        ws.cell(row, 1, name)
        ws.cell(row, 2, contract[field])
        set_name(wb, name, abs_ref(lib.SHEET_META, "B%d" % row))
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 12


def build_archetypes(wb, archetypes, styles):
    ws = recreate_sheet(wb, lib.SHEET_ARCH, len(wb.sheetnames))
    widths = {"A": 26, "B": 40, "C": 80}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    ws["A1"] = "Architecture Archetypes — reference (read-only)"
    copy_style(styles["title"], ws["A1"])
    ws.merge_cells("A1:C1")
    ws.row_dimensions[1].height = 25.5
    ws["A2"] = (
        "Plain-language description of each architecture pattern offered in the ARCHETYPE_PRIMARY / "
        "ARCHETYPE_SECONDARY questions of 'Extended Characterization'. Pick the one that best describes "
        "the application; if none fits, answer 'None of these fits' (the application then goes to "
        "Security review). This list is maintained centrally (contract/archetypes.yml) — do not edit it here."
    )
    copy_style(styles["note"], ws["A2"])
    ws.merge_cells("A2:C2")
    ws.row_dimensions[2].height = 48
    for col, header in zip("ABC", ("Archetype ID", "Name", "Description (plain language)")):
        ws[col + "4"] = header
        copy_style(styles["header"], ws[col + "4"])
    ws.row_dimensions[4].height = 21.75
    for row, a in enumerate(archetypes["archetypes"], start=5):
        for col, value in zip("ABC", (a["id"], a["name"], a["description"])):
            ws[col + str(row)] = value
            copy_style(styles["key"] if col == "A" else styles["question"], ws[col + str(row)])
        ws.row_dimensions[row].height = 30
    ws.freeze_panes = "A5"
    ws.protection.sheet = True


def build_extended(wb, characterization, contract, styles):
    keys = lib.keys_for_sheet(characterization, lib.SHEET_EXT)
    ws = recreate_sheet(wb, lib.SHEET_EXT, wb.sheetnames.index(lib.SHEET_CF) + 1)
    widths = {"A": 30, "B": 50, "C": 34, "D": 30, "E": 40, "F": 11, "G": 34, "H": 12, "I": 11}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w

    ws["A1"] = "Extended Characterization"
    copy_style(styles["title"], ws["A1"])
    ws.merge_cells("A1:I1")
    ws.row_dimensions[1].height = 25.5
    ws["A2"] = (
        "Fill in the 'Answer' column (yellow) for every row whose 'Applicable' column says Applicable; "
        "greyed-out rows don't apply given earlier answers and can stay blank. When done, the 'Status' "
        "column must show no 'Missing'. 'Don't know' is allowed, but on a mandatory question it sets "
        "Status to 'Review' and sends the application to Security review. These answers complement the "
        "'Characterization Form' (not repeated here) and feed the threat model. See the 'Archetypes' "
        "sheet for the architecture patterns."
    )
    copy_style(styles["note"], ws["A2"])
    ws.merge_cells("A2:I2")
    ws.row_dimensions[2].height = 48

    headers = contract_columns(contract, lib.SHEET_EXT)
    for col, header in headers.items():
        ws[col + "4"] = header
        copy_style(styles["header"], ws[col + "4"])
    ws.row_dimensions[4].height = 21.75

    # Pre-pass: assign rows, so conditions can reference any key's Applicable cell.
    rows, row = {}, 5
    layout = []
    for section in characterization["sections"]:
        section_keys = [k for k in keys if k["section"] == section["id"]]
        if not section_keys:
            continue
        layout.append(("section", section, row))
        row += 1
        for key in section_keys:
            rows[key["key"]] = row
            layout.append(("key", key, row))
            row += 1
    last = row - 1
    guarded = {k["key"] for k in keys if k.get("shown_when")}

    def ref(name):
        guard = '$H$%d="Applicable"' % rows[name] if name in guarded else None
        return name, guard

    for kind, item, r in layout:
        if kind == "section":
            ws["A%d" % r] = "%s. %s" % (item["id"], item["title"])
            for col in "ABCDEFGHI":
                cell = ws["%s%d" % (col, r)]
                cell.fill = SECTION_FILL
                cell.font = SECTION_FONT
                cell.alignment = Alignment(vertical="center")
                cell.border = copy.copy(styles["key"].border)
            ws.row_dimensions[r].height = 18
            continue
        key = item
        cells = {
            "A": (key["key"], styles["key"]),
            "B": (key["question"], styles["question"]),
            "C": (lib.options_label(key), styles["options"]),
            "D": (None, styles["answer"]),
            "E": (key.get("help"), styles["help"]),
            "F": ("Yes" if key["mandatory"] else "No", styles["options"]),
            "G": (lib.condition_label(key["shown_when"]) if key.get("shown_when") else None, styles["help"]),
            "H": (applicable_formula(key, ref), styles["computed"]),
            "I": (status_formula(r), styles["computed"]),
        }
        for col, (value, style) in cells.items():
            cell = ws["%s%d" % (col, r)]
            cell.value = value
            copy_style(style, cell)
        ws["F%d" % r].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        answer = ws["D%d" % r]
        answer.alignment = Alignment(vertical="center", wrap_text=True)
        if key["type"] == "date":
            answer.number_format = "yyyy-mm-dd"
        ws.row_dimensions[r].height = row_height(ws, r, widths, 60 if key["type"] == "long_text" else 45.75)
        set_name(wb, key["key"], abs_ref(lib.SHEET_EXT, "D%d" % r))
        add_answer_validation(ws, key, "D%d" % r)

    # Result block.
    r = last + 2
    ws["A%d" % r] = "Result (computed — do not edit)"
    for col in "ABCDEFGHI":
        cell = ws["%s%d" % (col, r)]
        cell.fill = SECTION_FILL
        cell.font = SECTION_FONT
        cell.border = copy.copy(styles["key"].border)
    ws.row_dimensions[r].height = 18
    special = {"status_any": lambda v: 'COUNTIF($I$5:$I$%d,"%s")>0' % (last, v)}
    triggers = [
        (lib.condition_to_excel(t["when"], ref, special), t["reason"]) for t in contract["review_triggers"]
    ]
    raw = "&".join('IF(%s,"%s; ","")' % (expr, reason) for expr, reason in triggers)
    missing_row, review_row, reasons_row = r + 1, r + 2, r + 3
    result = [
        (missing_row, "EXT_MISSING_COUNT", "Number of mandatory, applicable questions still unanswered.",
         '=COUNTIF($I$5:$I$%d,"Missing")' % last),
        (review_row, "NEEDS_SECURITY_REVIEW",
         "Yes when any Security review trigger fires (see governance document, 'Security review triggers').",
         '=IF(OR(%s),"Yes","No")' % ",".join(expr for expr, _ in triggers)),
        (reasons_row, "REVIEW_REASONS", "Why Security review is needed ('None' when it isn't).",
         '=IF($D$%d="Yes",LEFT(%s,LEN(%s)-2),"None")' % (review_row, raw, raw)),
    ]
    for rr, name, description, formula in result:
        ws["A%d" % rr] = name
        copy_style(styles["key"], ws["A%d" % rr])
        ws["B%d" % rr] = description
        copy_style(styles["question"], ws["B%d" % rr])
        ws["C%d" % rr] = None
        copy_style(styles["options"], ws["C%d" % rr])
        ws["D%d" % rr] = formula
        copy_style(styles["computed"], ws["D%d" % rr])
        ws.row_dimensions[rr].height = 45.75
        set_name(wb, name, abs_ref(lib.SHEET_EXT, "D%d" % rr))
    ws.merge_cells("D%d:I%d" % (reasons_row, reasons_row))

    # Conditional formatting.
    ws.conditional_formatting.add(
        "I5:I%d" % last, CellIsRule(operator="equal", formula=['"Missing"'], fill=RED_FILL, font=Font(color="FF9C0006"))
    )
    ws.conditional_formatting.add(
        "I5:I%d" % last, CellIsRule(operator="equal", formula=['"Review"'], fill=AMBER_FILL, font=Font(color="FF9C5700"))
    )
    ws.conditional_formatting.add(
        "A5:I%d" % last, FormulaRule(formula=['$H5="N/A"'], fill=GREY_FILL, font=Font(color="FF808080"))
    )
    ws.conditional_formatting.add(
        "D%d" % missing_row, CellIsRule(operator="greaterThan", formula=["0"], fill=RED_FILL, font=Font(color="FF9C0006"))
    )
    ws.conditional_formatting.add(
        "D%d" % review_row, CellIsRule(operator="equal", formula=['"Yes"'], fill=AMBER_FILL, font=Font(color="FF9C5700"))
    )
    ws.freeze_panes = "A5"
    return rows


def row_height(ws, r, widths, minimum):
    """Tall enough for the longest wrapped text in the row (about 1.15 characters per width unit)."""
    lines = 1
    for col in "BCEG":
        text = ws["%s%d" % (col, r)].value
        if text:
            lines = max(lines, -(-len(str(text)) // int(widths[col] * 1.15)))
    return max(minimum, 13 * lines + 8)


def applicable_formula(key, ref):
    if not key.get("shown_when"):
        return "Applicable"
    return '=IF(%s,"Applicable","N/A")' % lib.condition_to_excel(key["shown_when"], ref)


def status_formula(r):
    return (
        '=IF(H{r}="N/A","N/A",IF(D{r}="",IF(F{r}="Yes","Missing","Optional"),'
        'IF(AND(F{r}="Yes",D{r}="Don\'t know"),"Review","OK")))'
    ).format(r=r)


def add_answer_validation(ws, key, coordinate):
    if key["type"] == "choice":
        ws.add_data_validation(list_validation(key["list"], coordinate))
    elif key["type"] == "date":
        dv = DataValidation(type="date", operator="greaterThan", formula1="36526", allow_blank=True)
        dv.prompt = "Enter a date (YYYY-MM-DD)."
        dv.showInputMessage = True
        dv.error = "Enter a valid date (YYYY-MM-DD)."
        dv.showErrorMessage = True
        dv.add(coordinate)
        ws.add_data_validation(dv)


def contract_columns(contract, sheet):
    for s in contract["sheets"]:
        if s["name"] == sheet:
            return s["columns"]
    raise KeyError(sheet)


def add_remediation_owner(ws, last_row, vocabulary):
    ws["Q4"] = "Remediation Owner"
    copy_style(ws["P4"], ws["Q4"])
    comment = Comment(REMEDIATION_HELP, "Baseline")
    comment.width, comment.height = 380, 190
    ws["Q4"].comment = comment
    for r in range(5, last_row + 1):
        copy_style(ws["N%d" % r], ws["Q%d" % r])
    ws.column_dimensions["Q"].width = 18
    remove_validations_on(ws, "Q5")
    ws.add_data_validation(list_validation('"%s"' % ",".join(vocabulary), "Q5:Q%d" % last_row))
    ws.auto_filter.ref = "A4:Q%d" % last_row


def extend_custom_controls(ws):
    template_row = CC_V1_LAST_ROW
    for r in range(CC_V1_LAST_ROW + 1, CC_LAST_ROW + 1):
        for c in range(1, 17):
            copy_style(ws.cell(template_row, c), ws.cell(r, c))
        ws["G%d" % r] = "=TRUE()"
        ws["H%d" % r] = '=IF($A{r}="","",IF(AND(G{r},D{r}<=TARGET_LEVEL_NUM),"Applicable","N/A"))'.format(r=r)
        ws["P%d" % r] = (
            '=IF($A{r}="","",IF(H{r}="N/A","N/A",IFERROR(INDEX(CRIT_MATRIX,MATCH(O{r},CRIT_MATRIX_ROWS,0),'
            "MATCH(APP_RISK_TIER,CRIT_MATRIX_COLS,0)),\"\")))"
        ).format(r=r)
        ws.row_dimensions[r].height = ws.row_dimensions[template_row].height

    for dv in ws.data_validations.dataValidation:
        ranges = []
        for cr in dv.sqref.ranges:
            if cr.min_row == CC_FIRST_ROW and cr.max_row in (CC_V1_LAST_ROW, CC_LAST_ROW):
                ranges.append("%s%d:%s%d" % (get_column_letter(cr.min_col), CC_FIRST_ROW,
                                             get_column_letter(cr.max_col), CC_LAST_ROW))
            else:
                ranges.append(cr.coord)
        dv.sqref = MultiCellRange(" ".join(ranges))

    rebuilt = ConditionalFormattingList()
    for cf_range in ws.conditional_formatting:
        coords = []
        for cr in cf_range.sqref.ranges:
            if cr.min_row == CC_FIRST_ROW and cr.max_row in (CC_V1_LAST_ROW, CC_LAST_ROW):
                cr = "%s%d:%s%d" % (get_column_letter(cr.min_col), CC_FIRST_ROW,
                                    get_column_letter(cr.max_col), CC_LAST_ROW)
            coords.append(str(cr))
        for rule in cf_range.rules:
            rebuilt.add(" ".join(coords), rule)
    ws.conditional_formatting = rebuilt

    for merged in list(ws.merged_cells.ranges):
        if merged.min_row == 2 and merged.max_row == 2:
            ws.unmerge_cells(merged.coord)
    ws.merge_cells("A2:Q2")
    ws["A2"] = (
        "Use this sheet for requirements that matter to your organization but aren't part of the OWASP ASVS "
        "345-requirement set (e.g., an internal policy, a contractual/regulatory obligation specific to your "
        "company, or a control for a technology ASVS doesn't cover). Never add these rows to the 'Master "
        "Template' sheet — that sheet must stay a verbatim copy of ASVS so it can be safely updated from "
        "future ASVS releases. Column G here is a live formula (unlike the hidden helper column on Master "
        "Template) — edit it directly to control when a row applies; leave it as =TRUE() to make a row always "
        "applicable regardless of the application's characterization answers. It may reference any "
        "characterization key, including those on 'Extended Characterization' (e.g. SOURCING_MODEL).\n"
        "Control IDs: CUSTOM-NN for organization policy / contractual / platform-specific controls; "
        "TMX-<AREA>-NNN (e.g. TMX-WAF-001, TMX-VENDOR-001) for controls originated by threat modeling "
        "(COTS, vendor, compensating, infrastructure). Rows EXAMPLE-01 and TMX-EXAMPLE-001 are worked "
        "examples — edit or delete them."
    )
    ws.row_dimensions[2].height = 105


def write_tmx_example(ws):
    r = TMX_EXAMPLE_ROW
    values = {
        "A": "TMX-EXAMPLE-001",
        "B": "Vendor access (threat modeling)",
        "C": "Remote access",
        "D": 1,
        "E": (
            "EXAMPLE (edit or delete this row) — Vendor remote access to the environment goes only through "
            "the company PAM solution, is approved per session, and every session is recorded and retained."
        ),
        "F": "When the vendor has remote access (VENDOR_REMOTE_ACCESS is a 'Yes' option)",
        "G": (
            '=OR(VENDOR_REMOTE_ACCESS="Yes, on demand and controlled (approval, PAM, session recording)",'
            'VENDOR_REMOTE_ACCESS="Yes, permanent")'
        ),
        "I": "PAM configuration review + session log audit",
        "J": "Manual, periodic",
        "K": "Periodic review",
        "L": "Not assessed",
        "N": (
            "Worked example of a threat-modeling extension (TMX) control. Real TMX-* controls come from the "
            "threat modeling catalog — replace or delete this row."
        ),
        "O": "High",
    }
    for c in range(1, 18):
        col = get_column_letter(c)
        copy_style(ws["%s%d" % (col, CC_FIRST_ROW)], ws["%s%d" % (col, r)])
        if col in values:
            ws["%s%d" % (col, r)] = values[col]
        elif col not in ("H", "P"):
            ws["%s%d" % (col, r)] = None


INSTRUCTIONS = [
    ("title", "ASVS-Based Application Hardening Master Template — User Guide"),
    ("text", "Template v{template_version} — contract {contract_version} (OWASP ASVS {asvs_version})"),
    ("blank", None),
    ("head", "What this file is"),
    ("text", "This workbook is the master application-hardening baseline, built from the complete OWASP ASVS "
             "5.0.0 requirement set (345 requirements, 17 chapters, levels L1/L2/L3). It complements the "
             "governance document 'Application Hardening — Agnostic Baseline and Coverage Matrix'. It is "
             "maintained once, centrally, and then filtered per application without being rewritten each time. "
             "The same file is also the single intake for the application's threat model — one workbook per "
             "application serves both."),
    ("text", "Source — OWASP Application Security Verification Standard (ASVS) v5.0.0 — https://github.com/OWASP/ASVS"),
    ("head", "How to use it for a specific application"),
    ("text", "Step 1 — Make a copy of this file named after the application (e.g., Hardening_AppX.xlsx)."),
    ("text", "Step 2 — On the 'Characterization Form' sheet, fill in the 'Answer' column (highlighted in yellow) "
             "describing the nature of the application (target ASVS level, whether it exposes an API, uses "
             "OAuth/OIDC, is session-based, etc.). Use only the values offered in each dropdown — the closed "
             "vocabulary is what drives automatic filtering."),
    ("text", "Step 3 — On the 'Extended Characterization' sheet, fill in every yellow 'Answer' cell whose row "
             "is 'Applicable' (column H). Greyed-out rows don't apply given your earlier answers and can stay "
             "blank. When done, column I (Status) must show no 'Missing' (the count is at the bottom of the "
             "sheet). 'Don't know' is allowed, but on a mandatory question it shows 'Review'. For the "
             "architecture questions, the 'Archetypes' sheet describes each pattern in plain language."),
    ("text", "Step 4 — Check the result block at the bottom of 'Extended Characterization': "
             "NEEDS_SECURITY_REVIEW (Yes/No) and REVIEW_REASONS. If it says Yes, involve the Security team — "
             "the application has something the automatic threat model can't judge on its own."),
    ("text", "Step 5 — Run the threat model. Once both characterization sheets are complete (about 20-30 "
             "minutes in total), the threat model can already run from this file — it does not wait for the "
             "checklist below. It returns the likely threats, a detection plan, and a 'controls to verify "
             "first' list."),
    ("text", "Step 6 — Go to the 'Master Template' sheet. Column H (Applicable?) recalculates automatically from "
             "two things: (a) whether the requirement's ASVS chapter applies to this kind of application, and "
             "(b) whether the requirement's level is within the target level chosen in Step 2. Column P "
             "(Criticality) recalculates alongside it, from column O (Requirement Impact, fixed per ASVS "
             "chapter) crossed against this application's risk tier (computed from the Characterization Form's "
             "Data Sensitivity, Exposure, and Business Impact answers). Use the header AutoFilter to filter "
             "'Applicable' and work only with the requirements that actually apply to this application."),
    ("text", "Step 7 — For every requirement marked 'Applicable', fill in column L (Coverage Status) with the "
             "real status (Full coverage, Partial coverage, Gap, Not assessed), column M with the actual "
             "tool/evidence used at this company, and column N with notes. Start with the rows on the threat "
             "model's 'controls to verify first' list, then go through the remaining applicable rows. Do the "
             "same for the applicable rows of 'Custom Controls'."),
    ("text", "Step 8 — For every row marked 'Partial coverage' or 'Gap', fill in column Q (Remediation Owner) — "
             "who can actually close it: Internal dev (code change by our developers), Configuration (a setting "
             "in the product, platform or infrastructure), Vendor (only the vendor can fix it), Compensating "
             "(fixed outside the product: WAF, network segmentation, PAM, monitoring) or Risk acceptance "
             "(formally accepted instead of fixed). For vendor (COTS) products, most application-level gaps "
             "are Vendor or Configuration; use Compensating when the fix lives outside the product."),
    ("text", "Step 9 — Requirements marked 'N/A' need no status — they remain on record as 'considered and "
             "excluded by criterion', which is itself the audit evidence that the item was not overlooked, only "
             "ruled out based on the answers given in the Characterization Form."),
    ("text", "Step 10 — Column P (Criticality: Low/Medium/High/Critical) tells you which applicable requirements "
             "matter most for this specific application — it is not fixed per ASVS requirement, it depends on "
             "both the requirement's inherent impact and this application's risk profile. Items marked Critical "
             "or High are subject to the company's governance rules for high-risk findings (see the governance "
             "document's 'Governance rules by criticality' section)."),
    ("text", "Step 11 — Run the threat model again with the completed file. With Coverage Status filled in, it "
             "links the remaining gaps to the threats they leave open and produces a prioritized backlog."),
    ("blank", None),
    ("head", "Formalizing the output per application"),
    ("text", "Confluence — Once filled in, this filtered sheet is the data source for the per-application "
             "hardening record. Given the volume (100+ applications), a Confluence page per application (built "
             "from a page template) is the recommended format over a Word document per app — it stays "
             "versioned, searchable, linkable and commentable without generating a hundred loose files. The "
             "page's Threat Model and Detection Plan sections are filled from the threat modeling engine's "
             "output. How to turn a filtered sheet into a Confluence page is covered separately."),
    ("head", "This file also feeds the threat model"),
    ("text", "The threat modeling engine (https://github.com/Zzanoni/app-threat-modeling) reads this workbook "
             "directly: both characterization sheets, and the Coverage Status and Remediation Owner of every "
             "control. The workbook's structure is published as a machine-readable contract (the contract/ "
             "folder of https://github.com/Zzanoni/app-hardening-baseline). Do not rename sheets, keys, named "
             "ranges or columns in an application copy — the engine would no longer be able to read it."),
    ("head", "Maintaining the baseline"),
    ("text", "Updating a requirement — When a requirement changes (e.g., ASVS releases a new version, or a "
             "regulatory requirement is added), edit only this master file. The change propagates to any "
             "application whose next copy/update is generated from it. The master file is changed only through "
             "the migration scripts in the repository (scripts/migrate_*.py), never by hand. To move a filled-in "
             "copy made from an older template version, use scripts/upgrade_app_workbook.py instead of retyping."),
    ("head", "Scope note"),
    ("text", "Applicability in this template is scored at the ASVS chapter level (17 chapters), not per "
             "individual requirement. Within an applicable chapter, all requirements at or below the "
             "application's target level are considered applicable. This keeps the model auditable and "
             "maintainable across 100+ applications, at the cost of some requirements inside an applicable "
             "chapter being marginally relevant to a given app — those should be marked 'N/A' manually in "
             "column L with a one-line justification in column N."),
    ("head", "Color conventions"),
    ("text", "Yellow — Input cell — fill in per application (Characterization Form, Extended Characterization)."),
    ("text", "Light grey — Row computed as N/A for the current application (Master Template and Custom Controls "
             "column H; whole row on Extended Characterization)."),
    ("text", "Red / Amber (Extended Characterization, column I) — Status: Missing (red), Review (amber)."),
    ("text", "Red / Yellow / Green (column L) — Coverage status: Gap (red), Partial coverage (yellow), Full "
             "coverage (green)."),
    ("text", "Red / Orange / Yellow / Grey (column P) — Criticality: Critical (red), High (orange), Medium "
             "(yellow), N/A (grey)."),
    ("blank", None),
    ("head", "Adding controls beyond ASVS"),
    ("text", "If your organization needs a requirement that isn't part of ASVS (an internal policy, a "
             "contractual obligation, a control for something ASVS doesn't cover), add it to the 'Custom "
             "Controls' sheet, not to 'Master Template'. Master Template must stay a verbatim copy of ASVS so it "
             "can always be safely refreshed from a future ASVS release without losing anything. 'Custom "
             "Controls' has the same columns and the same Applicable?/Criticality mechanics, pre-wired and ready "
             "to fill in — see its own header note for how column G (Applicability Formula) works there."),
    ("text", "Control IDs — CUSTOM-NN for organization policy, contractual or platform-specific controls; "
             "TMX-<AREA>-NNN (e.g., TMX-WAF-001, TMX-VENDOR-001) for controls originated by threat modeling "
             "(COTS, vendor, compensating, infrastructure). EXAMPLE-01 and TMX-EXAMPLE-001 are worked examples."),
]


INSTRUCTION_STYLES = {
    "title": (Font(name="Arial", size=14, bold=True, color="FFFFFFFF"),
              PatternFill("solid", fgColor="FF1F4E78"), Alignment(vertical="bottom"), 27.75),
    "head": (Font(name="Arial", size=10, bold=True, color="FF1F4E78"),
             PatternFill("solid", fgColor="FFD9E2F3"), Alignment(vertical="bottom"), 18),
    "text": (Font(name="Arial", size=10), PatternFill(), Alignment(vertical="top", wrap_text=True), None),
}


def build_instructions(wb, contract):
    """Rewrite the Instructions sheet from INSTRUCTIONS (same look as v1)."""
    ws = wb["Instructions"]
    ws.delete_rows(1, ws.max_row)
    for r in list(ws.row_dimensions):
        del ws.row_dimensions[r]
    for r, (kind, text) in enumerate(INSTRUCTIONS, start=1):
        if kind == "blank":
            continue
        cell = ws.cell(r, 1, text.format(**contract))
        font, fill, alignment, height = INSTRUCTION_STYLES[kind]
        cell.font, cell.fill, cell.alignment = copy.copy(font), copy.copy(fill), copy.copy(alignment)
        if height is None:
            lines = -(-len(cell.value) // 120)
            height = max(18, 14 * lines + 6)
        ws.row_dimensions[r].height = height


def capture_styles(wb):
    cf = wb[lib.SHEET_CF]
    return {
        "title": cf["A1"],
        "note": cf["A2"],
        "header": cf["A4"],
        "key": cf["A5"],
        "question": cf["B5"],
        "options": cf["C5"],
        "answer": cf["D5"],
        "help": cf["E5"],
        "computed": cf["D17"],
    }


def reorder_sheets(wb, contract):
    order = [s["name"] for s in contract["sheets"]]
    wb._sheets = sorted(wb._sheets, key=lambda ws: order.index(ws.title))
    wb.active = 0
    for ws in wb:
        ws.sheet_view.tabSelected = ws.title == order[0]


def migrate(src, dst):
    characterization, archetypes, contract = lib.load_contract()
    wb = load_workbook(src)
    check_v1_layout(wb, characterization)
    before = snapshot_v1(wb)

    set_name(wb, "APP_NAME", abs_ref(lib.SHEET_CF, "D5"))
    styles = capture_styles(wb)
    build_extended(wb, characterization, contract, styles)
    build_archetypes(wb, archetypes, styles)
    build_lists(wb, characterization, archetypes)
    build_meta(wb, contract)

    vocabulary = contract["vocabularies"]["remediation_owner"]
    add_remediation_owner(wb[lib.SHEET_MT], MT_LAST_ROW, vocabulary)
    cc = wb[lib.SHEET_CC]
    extend_custom_controls(cc)
    add_remediation_owner(cc, CC_LAST_ROW, vocabulary)
    write_tmx_example(cc)
    build_instructions(wb, contract)
    reorder_sheets(wb, contract)

    after = snapshot_v1(wb)
    changed = [k for k in before if before[k] != after[k]]
    if changed:
        sys.exit("Refusing to save: v1 content changed in %s" % changed)
    lib.save_normalized(wb, dst)


def main():
    parser = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    parser.add_argument("--in", dest="src", default=lib.TEMPLATE)
    parser.add_argument("--out", dest="dst", default=None)
    args = parser.parse_args()
    migrate(args.src, args.dst or args.src)
    print("Migrated %s -> %s" % (args.src, args.dst or args.src))


if __name__ == "__main__":
    main()
