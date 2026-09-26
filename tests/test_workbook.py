"""Workbook regression tests (spec v2, section 8).

Run: python -m unittest discover -s tests -v
Tests that need formula results recalculate through LibreOffice and are skipped without it.
tests/fixtures/template_v1.xlsx is the v1.x template exactly as released (before v2).
"""

import datetime
import os
import shutil
import sys
import tempfile
import unittest
import zipfile

from openpyxl import load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(os.path.dirname(HERE), "scripts"))

import baseline_lib as lib  # noqa: E402
import migrate_v2  # noqa: E402
import upgrade_app_workbook  # noqa: E402

V1_FIXTURE = os.path.join(HERE, "fixtures", "template_v1.xlsx")
HAS_LO = lib.find_soffice() is not None
CHARACTERIZATION, ARCHETYPES, CONTRACT = lib.load_contract()
EXT_KEYS = lib.keys_for_sheet(CHARACTERIZATION, lib.SHEET_EXT)

V1_ANSWERS = [
    {"APP_NAME": "Portal", "TARGET_LEVEL": "L2", "API_EXPOSED": "Yes", "SPA_FRONTEND": "Yes",
     "SESSION_STATEFUL": "No", "FILE_HANDLING": "Yes", "OAUTH_OIDC_USED": "Yes",
     "DATA_SENSITIVITY": "Confidential", "WEBRTC_USED": "No", "EXPOSURE": "Public internet",
     "BUSINESS_IMPACT": "High"},
    {"APP_NAME": "Batch", "TARGET_LEVEL": "L1", "API_EXPOSED": "No", "SPA_FRONTEND": "No",
     "SESSION_STATEFUL": "Yes", "FILE_HANDLING": "No", "OAUTH_OIDC_USED": "No",
     "DATA_SENSITIVITY": "Regulated", "WEBRTC_USED": "Yes", "EXPOSURE": "Internal only",
     "BUSINESS_IMPACT": "Low"},
]


def tmpdir(test):
    path = tempfile.mkdtemp(prefix="baseline-test-")
    test.addCleanup(shutil.rmtree, path, ignore_errors=True)
    return path


def package_parts(path):
    with zipfile.ZipFile(path) as z:
        return [(name, z.read(name)) for name in z.namelist()]


def write_answers(wb, answers):
    for key, value in answers.items():
        sheet, cell = lib.parse_ref(lib.defined_name_ref(wb, key))
        wb[sheet][cell] = value


def recalculated(test, wb):
    path = os.path.join(tmpdir(test), "wb.xlsx")
    wb.save(path)
    result = lib.recalculate(path)
    test.addCleanup(shutil.rmtree, result.rsplit("/out/", 1)[0], ignore_errors=True)
    test.assertEqual(lib.formula_errors(result), [])
    return load_workbook(result, data_only=True)


def named(wb, name):
    sheet, cell = lib.parse_ref(lib.defined_name_ref(wb, name))
    return wb[sheet][cell].value


def risk_tier(v1):
    scores = {"None": 0, "Internal": 1, "Confidential": 2, "Regulated": 3,
              "Internal only": 0, "Partner/B2B network": 1, "Public internet": 2,
              "Low": 0, "Medium": 1, "High": 2, "Critical": 3}
    worst = max(scores.get(v1.get(k), 0) for k in ("DATA_SENSITIVITY", "EXPOSURE", "BUSINESS_IMPACT"))
    return ["Low", "Medium", "High", "Critical"][worst]


def expected_extended(answers):
    """Independent Python model of Applicable / Status / review triggers."""
    applicable, status = {}, {}

    def value_of(key):
        if key == "APP_RISK_TIER":
            return risk_tier(answers)
        if applicable.get(key) is False:
            return ""
        return answers.get(key, "")

    for key in EXT_KEYS:
        name = key["key"]
        applicable[name] = lib.evaluate_condition(key.get("shown_when") or {}, value_of)
        answer = answers.get(name)
        if not applicable[name]:
            status[name] = "N/A"
        elif answer in (None, ""):
            status[name] = "Missing" if key["mandatory"] else "Optional"
        elif key["mandatory"] and answer == "Don't know":
            status[name] = "Review"
        else:
            status[name] = "OK"
    special = {"status_any": lambda v: v in status.values()}
    fired = [t["reason"] for t in CONTRACT["review_triggers"]
             if lib.evaluate_condition(t["when"], value_of, special)]
    return applicable, status, fired


def complete_answers(**overrides):
    """A fully answered, review-free baseline: every key gets its first 'safe' answer."""
    answers = dict(V1_ANSWERS[0])
    for key in EXT_KEYS:
        if key["type"] == "choice":
            options = [o for o in key["options"] if o not in ("Don't know", "None of these fits")]
            answers[key["key"]] = "No" if key["key"] == "ARCHETYPE_DEVIATION" else options[0]
        elif key["type"] == "date":
            answers[key["key"]] = datetime.datetime(2026, 9, 26)
        else:
            answers[key["key"]] = "Sample answer"
    answers.update(overrides)
    return answers


class MigrationTests(unittest.TestCase):
    def test_idempotent_and_reproducible(self):
        d = tmpdir(self)
        first, second = os.path.join(d, "first.xlsx"), os.path.join(d, "second.xlsx")
        migrate_v2.migrate(V1_FIXTURE, first)
        migrate_v2.migrate(first, second)
        with open(first, "rb") as a, open(second, "rb") as b:
            self.assertEqual(a.read(), b.read(), "running the migration twice changed the file")
        # The committed template must be exactly what the script produces. Compared per package
        # part (not raw bytes), since deflate output may differ between zlib builds.
        self.assertEqual(package_parts(first), package_parts(lib.TEMPLATE),
                         "committed template differs from a fresh v1 -> v2 migration")

    def test_v1_content_untouched(self):
        v1, v2 = load_workbook(V1_FIXTURE), load_workbook(lib.TEMPLATE)
        for sheet, max_col in ((lib.SHEET_MT, 16), (lib.SHEET_CF, 5)):
            for row in v1[sheet].iter_rows(min_row=1, max_row=v1[sheet].max_row, max_col=max_col):
                for c in row:
                    self.assertEqual(v2[sheet][c.coordinate].value, c.value, "%s!%s" % (sheet, c.coordinate))
        for name in migrate_v2.V1_NAMES:
            self.assertEqual(lib.defined_name_ref(v2, name), lib.defined_name_ref(v1, name), name)
        self.assertTrue(v2[lib.SHEET_MT].column_dimensions["G"].hidden)
        self.assertEqual(v2[lib.SHEET_MT].auto_filter.ref, "A4:Q349")
        self.assertEqual(v2[lib.SHEET_CC].auto_filter.ref, "A4:Q204")


@unittest.skipUnless(HAS_LO, "LibreOffice not available")
class V1BehaviourTests(unittest.TestCase):
    def test_same_results_as_v1(self):
        for answers in V1_ANSWERS:
            with self.subTest(app=answers["APP_NAME"]):
                old, new = load_workbook(V1_FIXTURE), load_workbook(lib.TEMPLATE)
                write_answers(old, {k: v for k, v in answers.items() if k != "APP_NAME"})
                write_answers(new, answers)
                old, new = recalculated(self, old), recalculated(self, new)
                for r in range(17, 23):
                    self.assertEqual(new[lib.SHEET_CF]["D%d" % r].value, old[lib.SHEET_CF]["D%d" % r].value)
                self.assertEqual(named(new, "APP_RISK_TIER"), risk_tier(answers))
                for r in range(5, 350):
                    for col in "HP":
                        cell = "%s%d" % (col, r)
                        self.assertEqual(new[lib.SHEET_MT][cell].value, old[lib.SHEET_MT][cell].value, cell)
                for r in [5] + list(range(7, 25)):
                    for col in "HP":
                        cell = "%s%d" % (col, r)
                        self.assertEqual(new[lib.SHEET_CC][cell].value, old[lib.SHEET_CC][cell].value, cell)


@unittest.skipUnless(HAS_LO, "LibreOffice not available")
class ExtendedCharacterizationTests(unittest.TestCase):
    def assert_scenario(self, answers, expected_reasons, tmx_applicable=None):
        wb = load_workbook(lib.TEMPLATE)
        write_answers(wb, answers)
        wb = recalculated(self, wb)
        applicable, status, fired = expected_extended(answers)
        self.assertEqual(fired, expected_reasons, "Python model disagrees with the scenario")
        ws = wb[lib.SHEET_EXT]
        for key in EXT_KEYS:
            r = int(lib.parse_ref(lib.defined_name_ref(wb, key["key"]))[1][1:])
            self.assertEqual(ws["H%d" % r].value, "Applicable" if applicable[key["key"]] else "N/A", key["key"])
            self.assertEqual(ws["I%d" % r].value, status[key["key"]], key["key"])
        self.assertEqual(named(wb, "EXT_MISSING_COUNT"), list(status.values()).count("Missing"))
        self.assertEqual(named(wb, "NEEDS_SECURITY_REVIEW"), "Yes" if expected_reasons else "No")
        self.assertEqual(named(wb, "REVIEW_REASONS"), "; ".join(expected_reasons) or "None")
        if tmx_applicable is not None:
            self.assertEqual(wb[lib.SHEET_CC]["H6"].value, "Applicable" if tmx_applicable else "N/A")
        return wb

    def test_blank_template(self):
        self.assert_scenario({}, [])

    def test_inhouse_public_web(self):
        answers = complete_answers(SOURCING_MODEL="In-house (we develop it)",
                                   ARCHETYPE_PRIMARY="Public web application / portal")
        self.assert_scenario(answers, [], tmx_applicable=False)

    def test_internal_api_with_unknowns(self):
        answers = complete_answers(
            SOURCING_MODEL="In-house (we develop it)", ARCHETYPE_PRIMARY="Internal API / service",
            ARCHETYPE_DEVIATION="Don't know", MULTI_TENANT="Don't know", EXPOSURE="Internal only",
            DATA_SENSITIVITY="Internal", BUSINESS_IMPACT="Medium", WAF_IN_FRONT="")
        self.assert_scenario(answers, [
            "Components or flows outside the selected archetype (or unknown)",
            "Mandatory question answered Don't know",
        ])

    def test_cots_permanent_vendor_access(self):
        answers = complete_answers(
            SOURCING_MODEL="COTS (vendor product on our infrastructure)",
            ARCHETYPE_PRIMARY="Vendor product hosted by us", VENDOR_REMOTE_ACCESS="Yes, permanent",
            VENDOR_SUPPORTED="No", CUSTOMIZATIONS="Yes", CUSTOMIZATIONS_IN_REPO="",
            DATA_SENSITIVITY="Regulated")
        wb = self.assert_scenario(answers, [
            "Application Risk Tier is Critical",
            "Permanent or unknown vendor remote access",
            "Vendor version unsupported or support unknown",
        ], tmx_applicable=True)
        self.assertEqual(named(wb, "EXT_MISSING_COUNT"), 1)  # CUSTOMIZATIONS_IN_REPO

    def test_saas_and_stale_vendor_answers(self):
        # Vendor answers left over from a previous sourcing choice must not fire triggers.
        answers = complete_answers(
            SOURCING_MODEL="SaaS (vendor infrastructure)", ARCHETYPE_PRIMARY="None of these fits",
            VENDOR_REMOTE_ACCESS="Yes, permanent", VENDOR_SUPPORTED="Don't know",
            CUSTOMIZATIONS="Yes", HOSTING="")
        self.assert_scenario(answers, [
            "No archetype fits",
            "SaaS sourcing model (outside v1 archetypes)",
        ])


@unittest.skipUnless(HAS_LO, "LibreOffice not available")
class UpgradeTests(unittest.TestCase):
    def test_upgrade_v1_app_copy(self):
        d = tmpdir(self)
        app = load_workbook(V1_FIXTURE)
        answers = V1_ANSWERS[0]
        cf = app[lib.SHEET_CF]
        for r in range(5, 16):
            if cf["A%d" % r].value in answers:
                cf["D%d" % r] = answers[cf["A%d" % r].value]
        mt = app[lib.SHEET_MT]
        statuses = {"V1.1.1": ("Gap", "none", "no canonicalization"), "V6.2.1": ("Full coverage", "SAST", None),
                    "V17.3.2": ("Partial coverage", "DAST", "manual")}
        for r in range(5, 350):
            if mt["A%d" % r].value in statuses:
                mt["L%d" % r], mt["M%d" % r], mt["N%d" % r] = statuses[mt["A%d" % r].value]
        cc = app[lib.SHEET_CC]
        cc["A5"] = "CUSTOM-01"
        cc["E5"] = "Admin panels behind SSO"
        cc["L5"] = "Gap"
        cc["A6"], cc["D6"], cc["E6"], cc["G6"], cc["O6"] = "CUSTOM-02", 1, "Never", '=API_EXPOSED="Yes"', "Low"
        source = os.path.join(d, "app_v1.xlsx")
        app.save(source)

        output = os.path.join(d, "app_v2.xlsx")
        report = upgrade_app_workbook.upgrade(source, output, lib.TEMPLATE, keep_examples=False)
        self.assertEqual(report["answers"], 11)
        self.assertEqual(report["warnings"], [])
        self.assertEqual(report["examples_kept"], [])

        wb = load_workbook(output)
        for key, value in answers.items():
            self.assertEqual(named(wb, key), value, key)
        mt = wb[lib.SHEET_MT]
        found = {mt["A%d" % r].value: (mt["L%d" % r].value, mt["M%d" % r].value, mt["N%d" % r].value)
                 for r in range(5, 350) if mt["A%d" % r].value in statuses}
        self.assertEqual(found, statuses)
        cc = wb[lib.SHEET_CC]
        self.assertEqual([cc["A%d" % r].value for r in range(5, 9)], ["CUSTOM-01", "CUSTOM-02", None, None])
        self.assertEqual(cc["G6"].value, '=API_EXPOSED="Yes"')
        self.assertEqual(cc["G7"].value, "=TRUE()")

        wb = recalculated(self, wb)
        self.assertEqual(wb[lib.SHEET_CC]["H6"].value, "Applicable")
        self.assertEqual(wb[lib.SHEET_CC]["P6"].value, "Low")
        self.assertIn(wb[lib.SHEET_CC]["H7"].value, (None, ""))  # empty row: blank result

    def test_keep_examples(self):
        d = tmpdir(self)
        output = os.path.join(d, "out.xlsx")
        report = upgrade_app_workbook.upgrade(V1_FIXTURE, output, lib.TEMPLATE, keep_examples=True)
        # The v1 template still has EXAMPLE-01, so it is copied as a source row; TMX example kept.
        self.assertEqual(report["examples_kept"], ["TMX-EXAMPLE-001"])
        cc = load_workbook(output)[lib.SHEET_CC]
        self.assertEqual([cc["A5"].value, cc["A6"].value], ["EXAMPLE-01", "TMX-EXAMPLE-001"])


if __name__ == "__main__":
    unittest.main()
